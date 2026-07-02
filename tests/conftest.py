"""Shared test fixtures for upload reliability tests."""
import io
import sys
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import fitz
import pytest

# ── Path ──────────────────────────────────────────────────────────────────────
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


# ── File fixtures ─────────────────────────────────────────────────────────────

@pytest.fixture
def pdf_3_pages() -> bytes:
    """Minimal valid 3-page PDF created in-memory with PyMuPDF."""
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i + 1}")
    return doc.tobytes()


@pytest.fixture
def pdf_1_page() -> bytes:
    """Single-page PDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Single page")
    return doc.tobytes()


@pytest.fixture
def pptx_bytes() -> bytes:
    """Minimal valid PPTX file using python-pptx if available, else raw PPTX bytes."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        for i in range(5):
            slide_layout = prs.slide_layouts[5]
            slide = prs.slides.add_slide(slide_layout)
            tf = slide.shapes.title
            if tf:
                tf.text = f"Slide {i + 1}"
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
    except ImportError:
        pass

    # Fallback: read from the known uploads fixture if it exists
    fixture_path = Path(__file__).parent.parent / "uploads" / "f2112a3a-1075-4e29-a190-e9647a9f31e4_SEZG-501-01.pptx"
    if fixture_path.exists():
        return fixture_path.read_bytes()

    pytest.skip("python-pptx not installed and no PPTX fixture found")


@pytest.fixture
def corrupted_pdf_bytes() -> bytes:
    """Bytes that look like a PDF header but are corrupted."""
    return b"%PDF-1.4\n%%EOF\n\x00CORRUPTED DATA HERE"


@pytest.fixture
def random_bytes() -> bytes:
    """Non-file random bytes — no valid magic header."""
    return bytes(range(256)) * 4


@pytest.fixture
def executable_bytes() -> bytes:
    """ELF magic bytes (Linux executable header)."""
    return b"\x7fELF\x02\x01\x01\x00" + b"\x00" * 248


# ── DB / auth mocks ───────────────────────────────────────────────────────────

@pytest.fixture
def mock_user():
    from uuid import uuid4
    user = MagicMock()
    user.id = uuid4()
    user.role = "USER"
    return user


@pytest.fixture
def mock_db():
    """Async SQLAlchemy session mock."""
    db = AsyncMock()
    db.execute = AsyncMock()
    db.commit = AsyncMock()
    db.add = MagicMock()
    return db
